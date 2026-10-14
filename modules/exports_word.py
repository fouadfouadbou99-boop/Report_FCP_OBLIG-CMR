from docx import Document

def export_word(commentaire):

    doc = Document()

    doc.add_heading(
        "Rapport Comité",
        level=1
    )

    doc.add_paragraph(commentaire)

    filename = "Rapport_FCP.docx"

    doc.save(filename)

    return filename
